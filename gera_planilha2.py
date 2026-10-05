import re
from io import BytesIO
from pathlib import Path

import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, column_index_from_string

st.set_page_config(page_title="Gerador de Planilhas", page_icon="📊", layout="wide")

ROW_COL = "Nome"  # nome interno da 1ª coluna (nome do paciente)

# Colunas pré-definidas conforme planilha de referência (PLAN_MAIRA.xlsx)
DEFAULT_COLUMNS = [
    "Atendimento",
    "Data de Nacimento",
    "Leito",
    "CID",
    "ONCO/CLINICO/CIRURGICO",
    "IT OS/ENFERMARIA/UTI",
    "DATA/HORA DA INTUBAÇÃO",
    "DATA/HORA DA EXTUBAÇÃO",
    "TIPO DE EXTUBAÇÃO",
    "DESFECHO DA PACIENTE",
    "DATA/HORA DO DESFECHO",
    "TEMPO DE INTUBAÇÃO",
    "DATA/HORA RE-IOT",
    "MOTIVO DA FALHA DA EXTUBAÇÃO",
    "DATA/HORA TQT",
]
MAX_COLS, MAX_NAME = 16383, 255
BAD_SHEET = set('\\/?*[]:')
NUM = re.compile(r"^-?(0|[1-9]\d*)([.,]\d+)?$")
FUNC_NAME = re.compile(r"[A-Za-zÀ-ÿ_][A-Za-zÀ-ÿ0-9_.]*(?=\()")
REF_TOKEN = re.compile(r"\$?[A-Za-z]{1,4}\$?[0-9]+(?!\()")

# ---------------------------------------------------------------- catálogo de fórmulas
FORMULA_CATALOG = {
    "Matemática": [
        ("SOMA", "=SOMA(núm1; [núm2]; ...)", "Soma os valores de um intervalo ou lista de números."),
        ("MÉDIA", "=MÉDIA(núm1; [núm2]; ...)", "Calcula a média aritmética dos valores."),
        ("ARRED", "=ARRED(núm; núm_dígitos)", "Arredonda um número para a quantidade de dígitos informada."),
        ("ABS", "=ABS(núm)", "Retorna o valor absoluto (sem sinal) de um número."),
        ("POTÊNCIA", "=POTÊNCIA(núm; potência)", "Eleva um número a uma potência."),
        ("RAIZ", "=RAIZ(núm)", "Calcula a raiz quadrada de um número."),
        ("MOD", "=MOD(núm; divisor)", "Retorna o resto da divisão entre dois números."),
        ("SOMARPRODUTO", "=SOMARPRODUTO(matriz1; matriz2)", "Multiplica itens correspondentes de matrizes e soma os resultados."),
    ],
    "Estatística": [
        ("MÁXIMO", "=MÁXIMO(núm1; [núm2]; ...)", "Retorna o maior valor de um conjunto."),
        ("MÍNIMO", "=MÍNIMO(núm1; [núm2]; ...)", "Retorna o menor valor de um conjunto."),
        ("CONT.NÚM", "=CONT.NÚM(valor1; [valor2]; ...)", "Conta quantas células contêm números."),
        ("CONT.SE", "=CONT.SE(intervalo; critério)", "Conta células que atendem a um critério."),
        ("CONT.SES", "=CONT.SES(intervalo1; critério1; ...)", "Conta células que atendem a múltiplos critérios."),
        ("MED", "=MED(núm1; [núm2]; ...)", "Retorna a mediana de um conjunto de números."),
        ("DESVPAD", "=DESVPAD(núm1; [núm2]; ...)", "Calcula o desvio padrão de uma amostra."),
    ],
    "Lógica": [
        ("SE", "=SE(teste_lógico; valor_se_verdadeiro; valor_se_falso)", "Retorna um valor se a condição for verdadeira, outro se for falsa."),
        ("E", "=E(lógico1; [lógico2]; ...)", "Retorna VERDADEIRO se todos os argumentos forem verdadeiros."),
        ("OU", "=OU(lógico1; [lógico2]; ...)", "Retorna VERDADEIRO se ao menos um argumento for verdadeiro."),
        ("SEERRO", "=SEERRO(valor; valor_se_erro)", "Retorna um valor alternativo quando a fórmula gera erro."),
        ("NÃO", "=NÃO(lógico)", "Inverte o valor lógico (VERDADEIRO/FALSO) do argumento."),
    ],
    "Texto": [
        ("CONCATENAR", "=CONCATENAR(texto1; [texto2]; ...)", "Junta vários textos em um só."),
        ("ESQUERDA", "=ESQUERDA(texto; núm_caract)", "Extrai caracteres a partir do início de um texto."),
        ("DIREITA", "=DIREITA(texto; núm_caract)", "Extrai caracteres a partir do final de um texto."),
        ("MAIÚSCULA", "=MAIÚSCULA(texto)", "Converte o texto para letras maiúsculas."),
        ("MINÚSCULA", "=MINÚSCULA(texto)", "Converte o texto para letras minúsculas."),
        ("TEXTO", "=TEXTO(valor; formato)", "Converte um valor em texto com um formato específico."),
        ("ARRUMAR", "=ARRUMAR(texto)", "Remove espaços extras de um texto."),
    ],
    "Procura e Referência": [
        ("PROCV", "=PROCV(valor_procurado; matriz_tabela; núm_índice_coluna; [exato])", "Procura um valor na primeira coluna de uma tabela e retorna um valor na mesma linha."),
        ("PROCX", "=PROCX(valor_procurado; matriz_procurada; matriz_retornada)", "Versão moderna do PROCV/PROCH, mais flexível para localizar valores."),
        ("ÍNDICE", "=ÍNDICE(matriz; núm_linha; [núm_coluna])", "Retorna o valor em uma posição específica de uma matriz."),
        ("CORRESP", "=CORRESP(valor_procurado; matriz_procurada; [tipo_correspondência])", "Retorna a posição relativa de um valor dentro de um intervalo."),
    ],
    "Data e Hora": [
        ("HOJE", "=HOJE()", "Retorna a data atual do sistema."),
        ("AGORA", "=AGORA()", "Retorna a data e hora atuais do sistema."),
        ("DATA", "=DATA(ano; mês; dia)", "Monta uma data a partir de ano, mês e dia."),
        ("DIAS", "=DIAS(data_final; data_inicial)", "Calcula o número de dias entre duas datas."),
    ],
    "Financeira": [
        ("VP", "=VP(taxa; nper; pgto; [vf]; [tipo])", "Calcula o valor presente de um investimento ou empréstimo."),
        ("VF", "=VF(taxa; nper; pgto; [vp]; [tipo])", "Calcula o valor futuro de um investimento."),
        ("PGTO", "=PGTO(taxa; nper; vp; [vf]; [tipo])", "Calcula o pagamento periódico de um empréstimo."),
        ("TAXA", "=TAXA(nper; pgto; vp; [vf]; [tipo])", "Calcula a taxa de juros periódica de um empréstimo ou investimento."),
    ],
}
ALL_FUNC_NAMES = {name.casefold() for cat in FORMULA_CATALOG.values() for name, _, _ in cat}


def parse_address(addr: str, df: pd.DataFrame, data_cols: list[str]):
    """Converte um endereço estilo Excel (ex.: 'B4') em (índice_linha, nome_coluna, erro)."""
    m = re.match(r"^([A-Za-z]{1,3})(\d+)$", addr)
    if not m:
        return None, None, 'Endereço inválido. Use o formato de célula do Excel, por exemplo "B2" (letra = coluna, número = linha).'
    col_letters, row_num = m.group(1).upper(), int(m.group(2))
    try:
        col_idx = column_index_from_string(col_letters)
    except ValueError:
        return None, None, f'Coluna "{col_letters}" inválida.'
    if col_idx == 1:
        return None, None, 'A coluna A é reservada para o nome das linhas. Escolha uma coluna de dados (B, C, ...).'
    if row_num < 2:
        return None, None, "A linha 1 é o cabeçalho da tabela; os dados começam na linha 2."
    data_idx = col_idx - 2
    if data_idx < 0 or data_idx >= len(data_cols):
        return None, None, f'Coluna "{col_letters}" fora do intervalo atual (até {get_column_letter(len(data_cols) + 1)}).'
    row_idx = row_num - 2
    if row_idx >= len(df):
        return None, None, f"Linha {row_num} fora do intervalo atual (até a linha {len(df) + 1})."
    return row_idx, data_cols[data_idx], None


def rc_to_ref(r: int, c: str, data_cols: list[str]) -> str:
    """Converte (linha, coluna) internos para o endereço estilo Excel (ex.: 'B4')."""
    col_idx = data_cols.index(c) + 2
    return f"{get_column_letter(col_idx)}{r + 2}"


def coerce(v: str):
    """Texto numérico vira número; fórmula (começa com '=') é preservada; preserva zeros à esquerda."""
    v = v.strip()
    if v.startswith("="):
        return v
    if NUM.match(v) and len(v) <= 15:
        return float(v.replace(",", ".")) if re.search(r"[.,]", v) else int(v)
    return v


def validate_cell_ref_token(tok: str) -> list[str]:
    """Verifica se um token no formato 'coluna+linha' é uma referência de célula válida no Excel."""
    m = re.match(r"^\$?([A-Za-z]{1,4})\$?([0-9]+)$", tok)
    if not m:
        return []
    col, row = m.groups()
    issues = []
    if len(col) > 3:
        issues.append(f'Referência "{tok}": coluna com mais de 3 letras não existe no Excel (máximo é XFD).')
    if not (1 <= int(row) <= 1_048_576):
        issues.append(f'Referência "{tok}": linha {row} fora do intervalo válido do Excel (1 a 1.048.576).')
    return issues


def validate_formula(expr: str) -> list[str]:
    """Valida uma fórmula antes de aplicá-la a uma célula. Retorna lista de problemas (vazia = ok)."""
    issues = []
    e = expr.strip()
    if not e:
        issues.append("A fórmula está vazia.")
        return issues
    if not e.startswith("="):
        issues.append('Toda fórmula deve começar com "=".')
    body = e[1:] if e.startswith("=") else e
    if body.count("(") != body.count(")"):
        issues.append(f"Parênteses desbalanceados: {body.count('(')} aberto(s) e {body.count(')')} fechado(s).")
    if re.search(r"\(\s*\)", body):
        issues.append("Há uma função com parênteses vazios — faltam argumentos.")
    if body.strip() == "":
        issues.append("Não há conteúdo após o sinal de igual.")
    for match in FUNC_NAME.finditer(body):
        name = match.group(0)
        if name.casefold() not in ALL_FUNC_NAMES and not re.match(r"^[A-Z]+\d+$", name, re.I):
            issues.append(f'Função "{name}" não está no catálogo — confira se o nome está correto (ela ainda pode ser válida no Excel).')
    if re.search(r"[;,]\s*[;,]", body):
        issues.append("Há separadores de argumento duplicados ou um argumento vazio entre eles.")
    for match in REF_TOKEN.finditer(body):
        issues.extend(validate_cell_ref_token(match.group(0)))
    return issues


def validate(df: pd.DataFrame, sheet: str) -> list[str]:
    errs = []
    cols = list(df.columns[1:])
    if not cols:
        errs.append("Adicione pelo menos uma coluna.")
    if df.empty:
        errs.append("Adicione pelo menos uma linha.")
    names = df[ROW_COL].str.strip()
    if (names == "").any():
        errs.append(f"Há {(names == '').sum()} linha(s) sem nome (primeira coluna da tabela).")
    dup = names[(names != "") & names.str.casefold().duplicated(keep=False)].unique()
    if len(dup):
        errs.append("Nomes de linha duplicados: " + ", ".join(map(str, dup)))
    if any(len(n) > MAX_NAME for n in names):
        errs.append(f"Nome de linha com mais de {MAX_NAME} caracteres.")
    s = sheet.strip()
    if not s:
        errs.append("O nome da aba não pode ficar vazio.")
    elif len(s) > 31:
        errs.append("O nome da aba deve ter no máximo 31 caracteres.")
    elif set(s) & BAD_SHEET:
        errs.append("Nome da aba contém caracteres inválidos: " + " ".join(sorted(set(s) & BAD_SHEET)))
    elif s.startswith("'") or s.endswith("'"):
        errs.append("O nome da aba não pode começar/terminar com apóstrofo.")
    for i, row in df.iterrows():
        for c in cols:
            v = str(row[c]).strip()
            if v.startswith("="):
                probs = validate_formula(v)
                blocking = [p for p in probs if "não está no catálogo" not in p]
                if blocking:
                    errs.append(f'Fórmula inválida em "{row[ROW_COL]}" / "{c}": ' + "; ".join(blocking))
    return errs


def build_xlsx(df: pd.DataFrame, sheet: str, corner: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet.strip()
    hf, hfill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="FF4B4B")
    sfill = PatternFill("solid", fgColor="F0F2F6")
    ffill = PatternFill("solid", fgColor="E7F3E8")
    t = Side(style="thin", color="999999")
    border = Border(left=t, right=t, top=t, bottom=t)

    def put(r, c, val, font=None, fill=None, align="center"):
        cell = ws.cell(row=r, column=c)
        is_formula = isinstance(val, str) and val.startswith("=")
        cell.value = None if val == "" else val
        if isinstance(val, str) and val and not is_formula:
            cell.data_type = "s"  # texto comum nunca vira fórmula acidental
        cell.border = border
        cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
        if font: cell.font = font
        if fill: cell.fill = fill
        if is_formula and fill is None:
            cell.fill = ffill

    heads = [corner.strip()] + list(df.columns[1:])
    for j, h in enumerate(heads, 1):
        put(1, j, h, hf, hfill, "left" if j == 1 else "center")
    for i, row in enumerate(df.itertuples(index=False), 2):
        put(i, 1, row[0].strip(), Font(bold=True), sfill, "left")
        for j, v in enumerate(row[1:], 2):
            put(i, j, coerce(v))
    for j, h in enumerate(heads, 1):
        longest = max([len(h)] + [len(str(x)) for x in df.iloc[:, j - 1]])
        ws.column_dimensions[get_column_letter(j)].width = min(50, max(14, longest + 4))
    ws.freeze_panes = "B2"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- estado
if "df" not in st.session_state:
    st.session_state.df = pd.DataFrame(
        {ROW_COL: [f"Paciente {i}" for i in range(1, 4)],
         **{c: [""] * 3 for c in DEFAULT_COLUMNS}}
    )
    st.session_state.ver = 0
if "custom_formulas" not in st.session_state:
    st.session_state.custom_formulas = {}  # nome -> {"formula": str, "desc": str}
if "selected_fields" not in st.session_state:
    st.session_state.selected_fields = []


def apply(new_df: pd.DataFrame, msg: str | None = None):
    st.session_state.df = new_df.reset_index(drop=True)
    st.session_state.ver += 1  # recria o editor com a nova estrutura
    if msg:
        st.session_state.last_action = msg
    st.rerun()


# ---------------------------------------------------------------- página
st.title("📊 Gerador de Planilhas Excel")
st.caption("① Defina colunas e linhas na barra lateral → ② preencha a tabela → ③ use fórmulas → ④ baixe o .xlsx")

if st.session_state.get("last_action"):
    st.success(st.session_state.pop("last_action"))

with st.expander("Como usar", expanded=False):
    st.markdown(
        "- **Células:** duplo clique para editar. A 1ª coluna contém o **nome de cada linha**.\n"
        "- **Adicionar linha:** clique na linha vazia **+** no fim da tabela, ou use a barra lateral.\n"
        "- **Remover linhas:** na barra lateral, seção *Linhas*, selecione uma ou mais e clique em remover. "
        "Também dá para marcar a caixa à esquerda da linha e clicar na lixeira.\n"
        "- **Colunas:** adicionar, renomear e remover pela barra lateral.\n"
        "- **Fórmulas:** use a aba **📐 Fórmulas** abaixo da tabela para buscar funções do Excel, "
        "montar fórmulas personalizadas e inserir diretamente numa célula. Toda fórmula deve começar com \"=\"."
    )

with st.expander("🧾 Campos da planilha (PLAN_MAIRA.xlsx) — escolha o que usar", expanded=not st.session_state.selected_fields):
    st.caption("Marque os campos que deseja usar como colunas da tabela. Sua escolha substitui as colunas atuais pelos campos selecionados (os dados já preenchidos nas colunas mantidas são preservados).")
    chosen = st.multiselect(
        "Campos disponíveis", DEFAULT_COLUMNS,
        default=st.session_state.selected_fields, key="field_picker",
    )
    if st.button("✅ Aplicar seleção de campos", type="primary"):
        if not chosen:
            st.error("Selecione ao menos um campo.")
        else:
            base = st.session_state.df
            new_df = base[[ROW_COL]].copy()
            for c in chosen:
                new_df[c] = base[c] if c in base.columns else ""
            st.session_state.selected_fields = chosen
            st.session_state.last_action = f"Campos atualizados: {', '.join(chosen)}."
            apply(new_df)

edited = (
    st.data_editor(
        st.session_state.df,
        num_rows="dynamic",
        hide_index=True,
        key=f"editor_{st.session_state.ver}",
        column_config={ROW_COL: st.column_config.TextColumn("Nome do paciente", required=True)},
    )
    .fillna("")
    .astype(str)
    .reset_index(drop=True)
)

# ---------------------------------------------------------------- aba de fórmulas
st.subheader("📐 Fórmulas")
tab_cat, tab_custom, tab_direct, tab_check = st.tabs(
    ["📚 Catálogo do Excel", "⭐ Minhas fórmulas", "📍 Inserir em célula", "✅ Validar fórmula"]
)

data_cols = list(edited.columns[1:])
row_options = list(edited.index)


def cell_picker(key_prefix):
    if not row_options or not data_cols:
        st.info("Crie ao menos uma linha e uma coluna na tabela para inserir uma fórmula.")
        return None, None
    max_ref = rc_to_ref(len(edited) - 1, data_cols[-1], data_cols)
    addr = st.text_input(
        f"Endereço da célula (ex.: B2 — válido de B2 a {max_ref})",
        key=f"{key_prefix}_addr", placeholder="B2",
    ).strip().upper()
    if not addr:
        return None, None
    r, c, err = parse_address(addr, edited, data_cols)
    if err:
        st.error(err)
        return None, None
    atual = edited.at[r, c]
    preview = "vazio" if atual == "" else f"`{atual}`"
    st.caption(f'📍 **{addr}** → linha "{edited.at[r, ROW_COL]}", coluna "{c}". Conteúdo atual: {preview}')
    return r, c


with tab_cat:
    busca = st.text_input("🔎 Buscar função", placeholder="Ex.: soma, SE, procurar valor...", key="busca_cat")
    termo = busca.strip().casefold()
    for categoria, funcs in FORMULA_CATALOG.items():
        filtradas = [
            f for f in funcs
            if not termo or termo in f[0].casefold() or termo in f[2].casefold()
        ]
        if not filtradas:
            continue
        with st.expander(f"{categoria} ({len(filtradas)})", expanded=bool(termo)):
            for nome, sintaxe, desc in filtradas:
                st.markdown(f"**{nome}** — `{sintaxe}`")
                st.caption(desc)
                col_f, col_b = st.columns([3, 1])
                exemplo = col_f.text_input(
                    "Fórmula a inserir (edite os argumentos antes de aplicar)",
                    value=sintaxe, key=f"cat_input_{categoria}_{nome}", label_visibility="collapsed",
                )
                if col_b.button("Inserir", key=f"cat_btn_{categoria}_{nome}"):
                    st.session_state[f"pending_insert"] = exemplo
                    st.session_state[f"pending_source"] = f"cat_{categoria}_{nome}"
            if st.session_state.get("pending_source", "").startswith(f"cat_{categoria}_"):
                pass
    if st.session_state.get("pending_insert"):
        st.divider()
        st.info(f'Fórmula selecionada: `{st.session_state["pending_insert"]}`. Escolha a célula de destino abaixo.')
        probs = validate_formula(st.session_state["pending_insert"])
        bloqueios = [p for p in probs if "não está no catálogo" not in p]
        for p in probs:
            (st.error if p in bloqueios else st.warning)(p)
        r, c = cell_picker("insert_cat")
        if r is not None and st.button("✅ Confirmar inserção na célula", type="primary", disabled=bool(bloqueios)):
            ref = rc_to_ref(r, c, data_cols)
            edited.at[r, c] = st.session_state["pending_insert"]
            st.session_state.pending_insert = None
            apply(edited, msg=f"Fórmula aplicada em {ref}.")

with tab_custom:
    st.caption("Crie fórmulas próprias, salve com um nome e reutilize quando quiser.")
    with st.form("form_custom_formula", clear_on_submit=False):
        nome_f = st.text_input("Nome da fórmula", placeholder="Ex.: Lucro líquido")
        formula_f = st.text_input("Fórmula (comece com '=')", placeholder="=SOMA(B2:B10)-C2")
        desc_f = st.text_input("Descrição (opcional)", placeholder="Para que serve essa fórmula")
        salvar = st.form_submit_button("💾 Salvar fórmula")
    if salvar:
        n = nome_f.strip()
        probs = validate_formula(formula_f)
        bloqueios = [p for p in probs if "não está no catálogo" not in p]
        if not n:
            st.error("Dê um nome para a fórmula.")
        elif bloqueios:
            for p in bloqueios:
                st.error(p)
        else:
            st.session_state.custom_formulas[n] = {"formula": formula_f.strip(), "desc": desc_f.strip()}
            for p in [x for x in probs if x not in bloqueios]:
                st.warning(p)
            st.success(f'Fórmula "{n}" salva.')

    if not st.session_state.custom_formulas:
        st.info("Nenhuma fórmula personalizada salva ainda.")
    else:
        st.divider()
        for n, info in list(st.session_state.custom_formulas.items()):
            with st.expander(f"⭐ {n}"):
                st.markdown(f"`{info['formula']}`")
                if info["desc"]:
                    st.caption(info["desc"])
                novo_valor = st.text_input("Editar fórmula", value=info["formula"], key=f"edit_{n}")
                r, c = cell_picker(f"custom_{n}")
                b1, b2, b3 = st.columns(3)
                if b1.button("Salvar edição", key=f"save_{n}"):
                    probs = validate_formula(novo_valor)
                    bloqueios = [p for p in probs if "não está no catálogo" not in p]
                    if bloqueios:
                        for p in bloqueios:
                            st.error(p)
                    else:
                        st.session_state.custom_formulas[n]["formula"] = novo_valor.strip()
                        st.success("Fórmula atualizada.")
                        st.rerun()
                if b2.button("Inserir na célula", key=f"ins_{n}", disabled=r is None):
                    probs = validate_formula(info["formula"])
                    bloqueios = [p for p in probs if "não está no catálogo" not in p]
                    if bloqueios:
                        for p in bloqueios:
                            st.error(p)
                    else:
                        ref = rc_to_ref(r, c, data_cols)
                        edited.at[r, c] = info["formula"]
                        apply(edited, msg=f'Fórmula "{n}" aplicada em {ref}.')
                if b3.button("🗑️ Excluir", key=f"del_{n}"):
                    del st.session_state.custom_formulas[n]
                    st.rerun()

with tab_direct:
    st.caption("Escolha a célula exata (linha + coluna) e digite a fórmula ou valor a aplicar. "
               "Somente essa célula será alterada — nenhuma outra é tocada.")
    r, c = cell_picker("direct")
    if r is not None:
        atual = edited.at[r, c]
        st.caption(f'Conteúdo atual: `{atual or "(vazio)"}`')
        nova = st.text_area(
            "Conteúdo a aplicar (fórmulas devem começar com '=')",
            value=atual if atual else "=",
            key="direct_formula",
        )
        is_formula = nova.strip().startswith("=")
        bloqueios = []
        if is_formula:
            probs = validate_formula(nova)
            bloqueios = [p for p in probs if "não está no catálogo" not in p]
            for p in probs:
                (st.error if p in bloqueios else st.warning)(p)
        b1, b2 = st.columns(2)
        if b1.button("✅ Aplicar nesta célula", type="primary", disabled=bool(bloqueios), key="direct_apply"):
            edited.at[r, c] = nova.strip()
            apply(edited)
        if b2.button("🧹 Limpar esta célula", key="direct_clear"):
            edited.at[r, c] = ""
            apply(edited)

with tab_check:
    st.caption("Cole qualquer fórmula para verificar se está completa antes de usá-la.")
    teste = st.text_area("Fórmula para validar", placeholder="=SE(B2>100; \"Alto\"; \"Baixo\")")
    if teste.strip():
        probs = validate_formula(teste)
        if not probs:
            st.success("Fórmula com sintaxe aparentemente correta.")
        for p in probs:
            if "não está no catálogo" in p:
                st.warning(p)
            else:
                st.error(p)

st.divider()

# ---------------------------------------------------------------- sidebar
with st.sidebar:
    st.header("⚙️ Estrutura")
    sheet = st.text_input("Nome da aba", "Planilha1", max_chars=31)
    corner = st.text_input("Texto da célula A1 (opcional)", "")

    st.subheader("Colunas")
    new_col = st.text_input("Nova coluna", key="new_col", placeholder="Ex.: Janeiro")
    if st.button("➕ Adicionar coluna", use_container_width=True):
        n = new_col.strip()
        if not n:
            st.error("O nome da coluna não pode ficar vazio.")
        elif n.casefold() in {c.casefold() for c in edited.columns}:
            st.error(f'A coluna "{n}" já existe (ou é um nome reservado).')
        elif len(n) > MAX_NAME or len(edited.columns) > MAX_COLS:
            st.error("Nome muito longo ou limite de colunas do Excel atingido.")
        else:
            edited[n] = ""
            apply(edited, msg=f'Coluna "{n}" adicionada.')

    cols = list(edited.columns[1:])
    if cols:
        target = st.selectbox("Coluna selecionada", cols)
        new_name = st.text_input("Novo nome para a coluna selecionada", key="ren")
        c1, c2 = st.columns(2)
        if c1.button("✏️ Renomear", use_container_width=True):
            n = new_name.strip()
            others = {c.casefold() for c in cols if c != target} | {ROW_COL.casefold()}
            if not n:
                st.error("Informe o novo nome.")
            elif n.casefold() in others:
                st.error(f'Já existe uma coluna "{n}".')
            else:
                apply(edited.rename(columns={target: n}), msg=f'Coluna renomeada para "{n}".')
        if c2.button("🗑️ Remover", use_container_width=True):
            apply(edited.drop(columns=[target]), msg=f'Coluna "{target}" removida.')

    st.subheader("Linhas")
    new_row = st.text_input("Novo paciente", key="new_row", placeholder="Ex.: Maria da Silva")
    if st.button("➕ Adicionar linha", use_container_width=True):
        n = new_row.strip()
        if not n:
            st.error("O nome da linha não pode ficar vazio.")
        elif n.casefold() in {x.strip().casefold() for x in edited[ROW_COL]}:
            st.error(f'A linha "{n}" já existe.')
        else:
            edited.loc[len(edited)] = [n] + [""] * (len(edited.columns) - 1)
            apply(edited, msg=f'Linha "{n}" adicionada.')

    # ---- remover linhas já existentes
    if edited.empty:
        st.caption("Não há linhas para remover.")
    else:
        to_del = st.multiselect(
            "Linhas a remover",
            options=list(edited.index),
            format_func=lambda i: f"{i + 1} — {edited.at[i, ROW_COL].strip() or '(sem nome)'}",
            key=f"rm_rows_{st.session_state.ver}_{len(edited)}",
            placeholder="Escolha uma ou mais linhas",
        )
        if st.button("🗑️ Remover linhas selecionadas", use_container_width=True,
                     disabled=not to_del, type="primary"):
            apply(edited.drop(index=to_del), msg=f"{len(to_del)} linha(s) removida(s).")

# ---------------------------------------------------------------- exportação
st.divider()
errs = validate(edited, sheet)
n_r, n_c = len(edited), len(edited.columns) - 1
st.caption(f"**{n_r}** linha(s) × **{n_c}** coluna(s) → intervalo final "
           f"`A1:{get_column_letter(n_c + 1)}{n_r + 1}`")

for e in errs:
    st.error(e)

if not errs:
    data = build_xlsx(edited, sheet, corner)
    fname = st.text_input("Nome do arquivo", "planilha.xlsx")
    fname = fname.strip() or "planilha.xlsx"
    if not fname.lower().endswith(".xlsx"):
        fname += ".xlsx"
    if re.search(r'[\\/:*?"<>|]', fname):
        st.error('O nome do arquivo não pode conter \\ / : * ? " < > |')
    else:
        st.download_button("⬇️ Baixar planilha (.xlsx)", data, file_name=fname, type="primary",
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        with st.expander("Salvar no servidor (opcional)"):
            path = st.text_input("Caminho completo", str(Path.home() / fname))
            if st.button("💾 Salvar neste caminho"):
                try:
                    p = Path(path).expanduser()
                    if p.suffix.lower() != ".xlsx":
                        p = p.with_suffix(".xlsx")
                    p.parent.mkdir(parents=True, exist_ok=True)
                    p.write_bytes(data)
                    st.success(f"Planilha salva em: {p}")
                except PermissionError:
                    st.error("Sem permissão para gravar neste local ou arquivo em uso.")
                except OSError as ex:
                    st.error(f"Não foi possível gravar o arquivo: {ex}")
